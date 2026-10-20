"""M8 — parecer pré-audiência e roteiro de bolso em Word (.docx).

Todo documento sai como MINUTA: só vai a cliente depois de revisado pelo advogado.
"""

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROTULO_ONUS = {"reclamante": "Reclamante", "reclamada": "Reclamada", "pericia": "Perícia", None: "—"}
AVISO_MINUTA = ("MINUTA GERADA COM APOIO DE IA — revisada_por_advogado: NÃO. "
                "Conferir regras, fundamentos e datas antes de qualquer uso.")


def _brl(valor):
    if valor is None:
        return "—"
    inteiro, centavos = f"{valor:,.2f}".split(".")
    return f"R$ {inteiro.replace(',', '.')},{centavos}"


def _novo_documento(paisagem=False):
    doc = Document()
    if paisagem:
        secao = doc.sections[0]
        secao.orientation = WD_ORIENT.LANDSCAPE
        secao.page_width, secao.page_height = secao.page_height, secao.page_width
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(AVISO_MINUTA)
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    return doc


def _tabela(doc, cabecalho, linhas):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    for i, titulo in enumerate(cabecalho):
        cel = t.rows[0].cells[i]
        cel.text = ""
        cel.paragraphs[0].add_run(titulo).bold = True
    for linha in linhas:
        cells = t.add_row().cells
        for i, valor in enumerate(linha):
            cells[i].text = str(valor)
    return t


def _lista(doc, itens, estilo="List Bullet"):
    for item in itens:
        doc.add_paragraph(item, style=estilo)


def _cabecalho_processo(doc, caso):
    proc, contrato = caso["processo"], caso["contrato"]
    reclamada = (proc.get("reclamada") or {}).get("nome", "—")
    _tabela(doc, ["Campo", "Dado"], [
        ["Processo", proc["numero"]],
        ["Vara", proc.get("vara", "—")],
        ["Rito", "Sumaríssimo" if proc["rito"] == "sumarissimo" else "Ordinário"],
        ["Ajuizamento", f"{proc['data_ajuizamento']:%d/%m/%Y}"],
        ["Audiência", f"{proc['data_audiencia']:%d/%m/%Y}"],
        ["Reclamada", reclamada],
        ["Contrato", f"{contrato['admissao']:%d/%m/%Y} a "
                     f"{contrato['saida']:%d/%m/%Y}" if contrato.get("saida") else
                     f"desde {contrato['admissao']:%d/%m/%Y}"],
        ["Lado", caso.get("lado", "reclamante").capitalize()],
    ])


def gerar_parecer(resultado, caminho):
    caso, datas, quadro, kit = (resultado["caso"], resultado["datas"],
                                resultado["quadro"], resultado["kit"])
    doc = _novo_documento(paisagem=True)
    doc.add_heading("Parecer pré-audiência", level=0)
    _cabecalho_processo(doc, caso)

    doc.add_heading("1. Alertas antes da audiência", level=1)
    alertas = list(datas["alertas"])
    alertas.append(f"Marco da prescrição quinquenal: {datas['marco_quinquenal']:%d/%m/%Y}.")
    alertas += [t["alerta"] + (f" (conferido em {t['conferido_em']})" if t.get("conferido_em")
                               else " (NÃO conferido)") for t in resultado["temas"]]
    alertas += [linha["pendencia"] for linha in quadro if linha.get("pendencia")]
    _lista(doc, alertas)

    doc.add_heading("2. Quadro de ônus da prova por pedido", level=1)
    rg = resultado["regra_geral"]
    doc.add_paragraph(f"{rg['fundamento']}: {rg['texto'].strip()}")
    linhas = []
    for linha in quadro:
        p = linha["pedido"]
        obs = datas["por_pedido"].get(p["id"], "")
        linhas.append([
            p["id"], linha["rotulo"], _brl(p.get("valor")), ROTULO_ONUS[linha["onus"]],
            linha["fundamento"] or "—", ", ".join(linha["meio_prova"]) or "—",
            linha["efeito"] + (f" {obs}" if obs else ""),
        ])
    _tabela(doc, ["#", "Pedido", "Valor", "Ônus", "Fundamento", "Meio de prova", "Efeito na instrução"],
            linhas)

    doc.add_heading("3. Defesa da reclamada (contestação ou prevista)", level=1)
    for linha in quadro:
        d, tese = linha["defesa"], linha["tese"]
        if not d:
            continue
        origem = "contestação" if d.get("origem") == "contestacao" else "prevista"
        texto = f"[{linha['pedido']['id']}] {linha['rotulo']} — {origem}: "
        if tese:
            texto += f"{tese['rotulo']}. Documento esperado: {tese['documento_esperado']}. Se vier: {tese['se_vier']}"
        else:
            texto += f"tese '{d.get('tese')}' fora do catálogo (rules/teses-defesa.yaml) — advogado avalia."
        if d.get("fonte"):
            texto += f" (fonte: {d['fonte']})"
        doc.add_paragraph(texto, style="List Bullet")

    doc.add_heading("4. Kit de instrução", level=1)
    lim = kit["limite_testemunhas"]
    doc.add_paragraph(f"Limite de testemunhas por parte: {lim['quantidade']} ({lim['fundamento']}).")
    doc.add_paragraph(kit["lembrete_preposto"])
    for bloco in kit["blocos"]:
        doc.add_heading(f"[{bloco['pedido_id']}] {bloco['rotulo']} — ônus: {ROTULO_ONUS[bloco['onus']]}", level=2)
        doc.add_paragraph(bloco["objetivo"])
        if bloco["sem_modelo"]:
            doc.add_paragraph("Sem perguntas-base em rules/instrucao.yaml para este pedido — preparar manualmente.")
        for titulo, chave in (("Depoimento do preposto (perguntas fechadas)", "preposto"),
                              ("Testemunha própria (perguntas abertas, em funil)", "testemunha_propria"),
                              ("Testemunha da reclamada (um fato por vez, confronto com documento)",
                               "testemunha_adversa")):
            if bloco[chave]:
                doc.add_paragraph(titulo).runs[0].bold = True
                _lista(doc, bloco[chave], "List Number")

    doc.add_heading("5. Requerimentos", level=1)
    _lista(doc, kit["requerimentos"] or ["Nenhum requerimento derivado das regras."])

    doc.add_heading("6. Contradita", level=1)
    _lista(doc, kit["contradita"]["gatilhos"])
    doc.add_paragraph("Perguntas de qualificação:").runs[0].bold = True
    _lista(doc, kit["contradita"]["perguntas_qualificacao"], "List Number")
    doc.add_paragraph(kit["contradita"]["resposta_sumula_357"].strip())

    doc.add_heading("7. Protestos", level=1)
    doc.add_paragraph(kit["lembrete_protestos"])
    for prot in kit["protestos"]:
        doc.add_paragraph(f"{prot['quando']}: {prot['texto']}", style="List Bullet")

    doc.add_heading("8. Versões das regras usadas", level=1)
    doc.add_paragraph(f"rules/onus.yaml versão {resultado['versao_regras']} — revisado por: "
                      f"{resultado['revisado_por']}.")
    doc.save(caminho)
    return caminho


def gerar_roteiro_bolso(resultado, caminho):
    """Versão curta para levar à audiência: ônus, requerimentos e perguntas-chave."""
    caso, quadro, kit = resultado["caso"], resultado["quadro"], resultado["kit"]
    doc = _novo_documento()
    doc.add_heading(f"Roteiro de bolso — {caso['processo']['numero']}", level=1)
    lim = kit["limite_testemunhas"]
    doc.add_paragraph(f"Audiência {caso['processo']['data_audiencia']:%d/%m/%Y} · "
                      f"até {lim['quantidade']} testemunhas ({lim['fundamento']}).")
    if resultado["datas"]["alertas"]:
        _lista(doc, resultado["datas"]["alertas"])
    _tabela(doc, ["#", "Pedido", "Ônus", "O que fazer"],
            [[l["pedido"]["id"], l["rotulo"], ROTULO_ONUS[l["onus"]], l["efeito"]] for l in quadro])
    doc.add_heading("Requerimentos", level=2)
    _lista(doc, kit["requerimentos"] or ["—"])
    doc.add_heading("Perguntas-chave ao preposto", level=2)
    for bloco in kit["blocos"]:
        for pergunta in bloco["preposto"][:2]:
            doc.add_paragraph(f"[{bloco['pedido_id']}] {pergunta}", style="List Bullet")
    doc.add_heading("Protesto (modelo)", level=2)
    doc.add_paragraph(kit["protestos"][0]["texto"])
    doc.save(caminho)
    return caminho
