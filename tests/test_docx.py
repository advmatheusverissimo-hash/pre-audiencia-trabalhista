import json

from docx import Document

from engine.cli import main


def _texto(caminho):
    doc = Document(caminho)
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            partes += [c.text for c in row.cells]
    return "\n".join(partes)


def test_cli_gera_parecer_e_roteiro(tmp_path, caso_bruto, capsys):
    entrada = tmp_path / "caso.json"
    entrada.write_text(json.dumps(caso_bruto), encoding="utf-8")
    assert main([str(entrada), "--saida", str(tmp_path / "out"), "--hoje", "2026-10-20"]) == 0
    arquivos = sorted((tmp_path / "out").glob("*.docx"))
    assert [a.name.split("-")[0] for a in arquivos] == ["parecer", "roteiro"]

    parecer = _texto(next(a for a in arquivos if a.name.startswith("parecer")))
    assert "revisada_por_advogado: NÃO" in parecer
    assert "Súmula 338, III, TST" in parecer
    assert "R$ 42.000,00" in parecer
    assert "Súmula 357" in parecer and "art. 795" in parecer.lower()
    assert "NÃO conferido" in parecer

    roteiro = _texto(next(a for a in arquivos if a.name.startswith("roteiro")))
    assert "até 3 testemunhas" in roteiro


def test_cli_caso_invalido_retorna_erro(tmp_path, caso_bruto, capsys):
    caso_bruto["processo"]["rito"] = "x"
    entrada = tmp_path / "caso.json"
    entrada.write_text(json.dumps(caso_bruto), encoding="utf-8")
    assert main([str(entrada), "--saida", str(tmp_path / "out")]) == 2
    assert "Erro:" in capsys.readouterr().err
    assert not (tmp_path / "out").exists()


def test_cli_arquivo_inexistente_e_json_quebrado(tmp_path, capsys):
    assert main([str(tmp_path / "nao-existe.json")]) == 2
    quebrado = tmp_path / "q.json"
    quebrado.write_text("{", encoding="utf-8")
    assert main([str(quebrado)]) == 2
    assert "JSON inválido" in capsys.readouterr().err


def test_cli_avisa_pedido_sem_regra(tmp_path, caso_bruto, capsys):
    caso_bruto["pedidos"].append({"id": "P9", "tipo": "pedido_inexistente", "descricao": "x"})
    entrada = tmp_path / "caso.json"
    entrada.write_text(json.dumps(caso_bruto), encoding="utf-8")
    assert main([str(entrada), "--saida", str(tmp_path / "out"), "--hoje", "2026-10-20"]) == 0
    assert "P9" in capsys.readouterr().err
